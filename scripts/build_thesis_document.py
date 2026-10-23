"""Build the formatted University of Debrecen thesis draft from Markdown chapters."""

from __future__ import annotations

import os
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
OUTPUT = ROOT / "Legal_Clause_Classification_Thesis_Draft.docx"
FONT = "Times New Roman"

CANDIDATE_NAME = os.getenv("THESIS_CANDIDATE_NAME", "[Insert candidate full name]")
PROGRAMME_NAME = os.getenv("THESIS_PROGRAMME_NAME", "Data Science MSc")
SUPERVISOR_NAME = os.getenv("THESIS_SUPERVISOR_NAME", "[Insert supervisor full name]")
SUPERVISOR_TITLE = os.getenv("THESIS_SUPERVISOR_TITLE", "[Insert academic title]")

CHAPTERS = [
    DOCS / "introduction.md",
    DOCS / "literature_review.md",
    DOCS / "methodology.md",
    DOCS / "results.md",
    DOCS / "discussion.md",
    DOCS / "limitations.md",
    DOCS / "conclusion.md",
    DOCS / "bibliography.md",
]

TITLE = (
    "An Empirical Comparison of Prompting, Parameter-Efficient Fine-Tuning, "
    "and Retrieval-Augmented Generation for Legal Contract Clause Classification"
)

ABSTRACT = (
    "This thesis compares prompting, parameter-efficient fine-tuning (PEFT), "
    "retrieval-augmented generation (RAG), and a PEFT-RAG hybrid for classifying "
    "contract clauses into 20 categories derived from the Contract Understanding "
    "Atticus Dataset, using TF-IDF logistic regression as a supervised baseline. "
    "All methods were assessed on grouped data partitions, and the principal "
    "analysis used accuracy, Macro F1, output-validity audits, class-level results, "
    "and cluster-bootstrap intervals over document-title groups. In the original "
    "five-method comparison, TF-IDF logistic regression produced the strongest "
    "result, with 0.8605 accuracy and 0.8131 Macro F1, while RAG and the hybrid "
    "outperformed zero-shot prompting and the original PEFT implementation. A "
    "post-hoc, validation-selected PEFT sensitivity analysis using answer-only "
    "supervision and matched training and inference formatting reached 0.8946 "
    "accuracy and 0.8460 Macro F1, but is reported separately because it was "
    "developed after inspection of the original PEFT behaviour. Performance "
    "remained substantially lower on identical clause texts associated with more "
    "than one label, revealing ambiguity that clause-only models cannot reliably "
    "resolve. The findings show that adaptation quality depends not only on the "
    "model family, but also on supervision design, prompt consistency, retrieval "
    "quality, output parsing, and the information preserved in the classification task."
)


def set_run_font(run, size: float | None = None, bold=None, italic=None):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)


def configure_page(section):
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.left_margin = Cm(3)
    section.right_margin = Cm(2)
    section.top_margin = Cm(3)
    section.bottom_margin = Cm(3)
    section.header_distance = Cm(1.25)
    section.footer_distance = Cm(1.25)


def set_page_number_start(section, value: int):
    sect_pr = section._sectPr
    pg_num = sect_pr.find(qn("w:pgNumType"))
    if pg_num is None:
        pg_num = OxmlElement("w:pgNumType")
        sect_pr.append(pg_num)
    pg_num.set(qn("w:start"), str(value))


def add_field(paragraph, instruction: str, display: str = ""):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = display
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    set_run_font(run, 10)


def add_page_number(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_field(p, " PAGE ", "1")


def keep_table_rows_together(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def shade_cell(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def remove_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders.append(el)


def set_cell_margins(cell, value: int = 90):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side in ("top", "start", "bottom", "end"):
        node = OxmlElement(f"w:{side}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        tc_mar.append(node)


def add_inline(paragraph, text: str):
    """Add a small Markdown inline subset: bold, italics, and code."""
    pattern = re.compile(r"(\*\*.+?\*\*|(?<!\*)\*[^*]+?\*(?!\*)|`[^`]+?`)")
    pos = 0
    for match in pattern.finditer(text):
        if match.start() > pos:
            set_run_font(paragraph.add_run(text[pos : match.start()]), 12)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, 12, bold=True)
        elif token.startswith("*"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, 12, italic=True)
        else:
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, 11)
        pos = match.end()
    if pos < len(text):
        set_run_font(paragraph.add_run(text[pos:]), 12)


def add_body_paragraph(doc, text: str, style: str = "Thesis Body"):
    p = doc.add_paragraph(style=style)
    add_inline(p, text)
    return p


def parse_table_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def add_markdown_table(doc, rows: list[list[str]]):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for i, values in enumerate(rows):
        row = table.rows[i]
        keep_table_rows_together(row)
        if i == 0:
            repeat_header(row)
        for j, value in enumerate(values):
            cell = row.cells[j]
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            if j > 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_inline(p, value.replace("\\|", "|"))
            for run in p.runs:
                run.font.size = Pt(9.5)
                if i == 0:
                    run.bold = True
            if i == 0:
                shade_cell(cell, "D9E2F3")
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def markdown_blocks(text: str):
    lines = text.splitlines()
    i = 0
    buffer: list[str] = []

    def flush():
        nonlocal buffer
        if buffer:
            yield ("paragraph", " ".join(s.strip() for s in buffer))
            buffer = []

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if not stripped:
            yield from flush()
            i += 1
            continue
        image = re.match(r"^!\[(.+?)\]\((.+?)\)$", stripped)
        if image:
            yield from flush()
            yield ("image", (image.group(1), image.group(2)))
            i += 1
            continue
        if (
            stripped.startswith("|")
            and i + 1 < len(lines)
            and re.match(r"^\|?\s*:?-+", lines[i + 1].strip())
        ):
            yield from flush()
            table_lines = [stripped]
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
            yield ("table", [parse_table_row(x) for x in table_lines])
            continue
        heading = re.match(r"^(#{1,3})\s+(.+)$", stripped)
        if heading:
            yield from flush()
            yield ("heading", (len(heading.group(1)), heading.group(2)))
            i += 1
            continue
        list_item = re.match(r"^([1-9]\.|[-*])\s+(.+)$", stripped)
        if list_item:
            yield from flush()
            marker, content = list_item.groups()
            continuation = []
            i += 1
            while i < len(lines):
                nxt = lines[i]
                if (
                    not nxt.strip()
                    or re.match(r"^(#{1,3})\s+", nxt.strip())
                    or re.match(r"^([1-9]\.|[-*])\s+", nxt.strip())
                    or re.match(r"^!\[.+?\]\(.+?\)$", nxt.strip())
                    or nxt.strip().startswith("|")
                ):
                    break
                continuation.append(nxt.strip())
                i += 1
            yield ("list", (marker, " ".join([content] + continuation)))
            continue
        buffer.append(stripped)
        i += 1
    yield from flush()


def add_markdown(doc, path: Path):
    for kind, payload in markdown_blocks(path.read_text(encoding="utf-8")):
        if kind == "heading":
            level, title = payload
            if level == 1:
                p = doc.add_paragraph(style="Heading 1")
                p.paragraph_format.page_break_before = True
                add_inline(p, title.replace(" - ", ": ", 1))
            else:
                p = doc.add_paragraph(style=f"Heading {level}")
                add_inline(p, title)
        elif kind == "table":
            add_markdown_table(doc, payload)
        elif kind == "list":
            marker, text = payload
            style = "List Number" if marker[0].isdigit() else "List Bullet"
            add_body_paragraph(doc, text, style=style)
        elif kind == "image":
            caption, relative_path = payload
            image_path = path.parent / relative_path
            if not image_path.exists():
                raise FileNotFoundError(image_path)
            picture = doc.add_paragraph()
            picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
            picture.paragraph_format.keep_with_next = True
            run = picture.add_run()
            run.add_picture(str(image_path), width=Cm(14.5))
            caption_paragraph = doc.add_paragraph(style="Caption")
            add_inline(caption_paragraph, caption)
        else:
            text = payload
            if re.match(r"^\*\*Table \d+\.\d+\..+\*\*$", text):
                p = doc.add_paragraph(style="Caption")
                add_inline(p, text[2:-2])
            else:
                add_body_paragraph(doc, text)


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)

    if "Thesis Body" not in styles:
        body = styles.add_style("Thesis Body", WD_STYLE_TYPE.PARAGRAPH)
    else:
        body = styles["Thesis Body"]
    body.base_style = normal
    body.font.name = FONT
    body._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    body._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    body.font.size = Pt(12)
    body.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    body.paragraph_format.line_spacing = 1.5
    body.paragraph_format.first_line_indent = Cm(1.25)
    body.paragraph_format.space_after = Pt(6)
    body.paragraph_format.widow_control = True

    declaration = styles.add_style("Declaration Body", WD_STYLE_TYPE.PARAGRAPH)
    declaration.base_style = normal
    declaration.font.name = FONT
    declaration._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    declaration._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    declaration.font.size = Pt(12)
    declaration.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    declaration.paragraph_format.line_spacing = 1.15
    declaration.paragraph_format.first_line_indent = Cm(0)
    declaration.paragraph_format.space_after = Pt(4)

    for name in ("List Number", "List Bullet"):
        style = styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style.font.size = Pt(12)
        style.paragraph_format.line_spacing = 1.5
        style.paragraph_format.space_after = Pt(3)

    for name, size, before, after in (
        ("Heading 1", 16, 0, 18),
        ("Heading 2", 14, 18, 8),
        ("Heading 3", 12, 12, 6),
    ):
        style = styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    cap = styles["Caption"]
    cap.font.name = FONT
    cap._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    cap._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    cap.font.size = Pt(11)
    cap.font.bold = True
    cap.font.italic = False
    cap.font.color.rgb = RGBColor(0, 0, 0)
    cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.keep_with_next = True
    cap.paragraph_format.space_before = Pt(8)
    cap.paragraph_format.space_after = Pt(5)


def add_title_page(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run("UNIVERSITY OF DEBRECEN"), 14, bold=True)
    p.paragraph_format.space_after = Pt(5)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run("FACULTY OF INFORMATICS"), 14, bold=True)

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(48)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run("THESIS"), 16, bold=True)
    p.paragraph_format.space_after = Pt(28)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1.15
    set_run_font(p.add_run(TITLE), 16, bold=True)

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(78)
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Cm(7.2)
    table.columns[1].width = Cm(7.2)
    remove_table_borders(table)
    left, right = table.rows[0].cells
    for cell in (left, right):
        set_cell_margins(cell, 80)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    p = left.paragraphs[0]
    set_run_font(p.add_run("Supervisor"), 12, bold=True)
    p = left.add_paragraph()
    set_run_font(p.add_run(SUPERVISOR_NAME), 12)
    p = left.add_paragraph()
    set_run_font(p.add_run(SUPERVISOR_TITLE), 12)
    p = right.paragraphs[0]
    set_run_font(p.add_run("Candidate"), 12, bold=True)
    p = right.add_paragraph()
    set_run_font(p.add_run(CANDIDATE_NAME), 12)
    p = right.add_paragraph()
    set_run_font(p.add_run(PROGRAMME_NAME), 12)

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(42)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run("DEBRECEN"), 12, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run("2026"), 12, bold=True)


def add_acknowledgements(doc):
    p = doc.add_paragraph(style="Heading 1")
    add_inline(p, "Acknowledgements")
    supervisor_reference = (
        SUPERVISOR_NAME
        if not SUPERVISOR_NAME.startswith("[Insert")
        else "my supervisor"
    )
    add_body_paragraph(
        doc,
        f"I would like to thank {supervisor_reference} for guidance, constructive "
        "feedback and support throughout the development of this thesis. The "
        "comments received helped me refine the research design and present the "
        "findings more critically.",
    )
    add_body_paragraph(
        doc,
        "I am also grateful to my family and friends for their encouragement, "
        "patience and practical support during my studies. Finally, I would like "
        "to thank everyone who contributed advice, discussion or moral support "
        "while I worked through the technical and academic challenges of this "
        "research.",
    )
    doc.add_page_break()


def add_front_matter(doc):
    p = doc.add_paragraph(style="Heading 1")
    add_inline(p, "Abstract")
    add_body_paragraph(doc, ABSTRACT)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    add_inline(
        p,
        "**Keywords:** legal NLP; contract clause classification; prompting; "
        "PEFT; retrieval-augmented generation",
    )

    doc.add_page_break()
    # A genuine TOC field replaces the earlier manually typed contents list.
    # Word refreshes the entries from the Heading 1--3 styles when the document
    # is opened (or when the user presses Ctrl+A, F9).
    p = doc.add_paragraph(style="Title")
    add_inline(p, "Table of Contents")
    toc = doc.add_paragraph()
    add_field(
        toc,
        ' TOC \\o "1-3" \\h \\z \\u ',
        "Update the table in Word by pressing Ctrl+A, then F9.",
    )


def add_ai_declaration(doc):
    p = doc.add_paragraph(style="Heading 1")
    p.paragraph_format.page_break_before = True
    add_inline(p, "Declaration on the Use of Artificial Intelligence Tools")
    add_body_paragraph(
        doc,
        "I hereby declare that during the preparation of my thesis I used "
        "artificial intelligence-based tools as described below.",
        style="Declaration Body",
    )
    p = doc.add_paragraph(style="Heading 2")
    add_inline(p, "1. Use of AI tools")
    add_body_paragraph(
        doc,
        "☐ I did not use any artificial intelligence-based tools during the "
        "preparation of my thesis.",
        style="Declaration Body",
    )
    add_body_paragraph(
        doc,
        "☒ I used artificial intelligence-based tools during the preparation "
        "of my thesis as follows:",
        style="Declaration Body",
    )
    table = doc.add_table(rows=2, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    entries = [
        (
            "Name and version of the AI tool",
            "OpenAI ChatGPT/Codex (versions available during the 2025–2026 research period)",
        ),
        (
            "Exactly how it was used",
            "Language correction and stylistic refinement; assistance with "
            "drafting and restructuring text; explanation and debugging of "
            "program code; checking consistency between saved experimental "
            "outputs, tables and narrative; and generating suggestions for "
            "literature-search terms. All experimental results, interpretations, "
            "citations and final wording were reviewed and verified by the author.",
        ),
    ]
    for i, (a, b) in enumerate(entries):
        for j, value in enumerate((a, b)):
            cell = table.cell(i, j)
            set_cell_margins(cell, 110)
            p = cell.paragraphs[0]
            p.paragraph_format.line_spacing = 1.15
            add_inline(p, value)
            if j == 0:
                for run in p.runs:
                    run.bold = True
                shade_cell(cell, "D9E2F3")

    p = doc.add_paragraph(style="Heading 2")
    add_inline(p, "2. Declaration of independent work")
    statements = [
        "I declare that I take full responsibility for the professional content "
        "of my thesis. I acknowledge that the use of artificial "
        "intelligence-based tools does not replace independent work; therefore, "
        "I have verified all statements, results, interpretations, conclusions, "
        "and references included in the thesis.",
        "I further declare that the final content of the thesis was compiled by "
        "me; any AI-generated or modified content has been reviewed, verified, "
        "and corrected where necessary; the original results, interpretations, "
        "and conclusions presented in the thesis are my responsibility; and the "
        "use of AI tools has been truthfully and fully disclosed in this declaration.",
    ]
    for statement in statements:
        add_body_paragraph(doc, statement, style="Declaration Body")
    p = doc.add_paragraph(style="Declaration Body")
    p.paragraph_format.space_before = Pt(12)
    add_inline(
        p,
        "Date: ____________________                         "
        "Signature: ______________________________",
    )


def build():
    doc = Document()
    configure_page(doc.sections[0])
    doc.sections[0].different_first_page_header_footer = True
    configure_styles(doc)
    doc.core_properties.title = TITLE
    doc.core_properties.subject = "MSc thesis on legal contract clause classification"
    doc.core_properties.author = (
        CANDIDATE_NAME if not CANDIDATE_NAME.startswith("[Insert") else ""
    )
    settings = doc.settings._element
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    settings.append(update)

    add_title_page(doc)
    main = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_page(main)
    main.header.is_linked_to_previous = False
    main.footer.is_linked_to_previous = False
    set_page_number_start(main, 1)
    add_page_number(main)
    add_acknowledgements(doc)
    add_front_matter(doc)

    for chapter in CHAPTERS:
        add_markdown(doc, chapter)
    add_ai_declaration(doc)

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
