"""Revise the user-edited thesis DOCX while preserving Chapters 1 and 2.

The script converts the manually typed contents list into a real Word TOC
field, applies the user's less hyphenated prose style to Chapters 3--7, and
replaces the bibliography with the verified repository bibliography.
"""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.text.paragraph import Paragraph


STYLE_REPLACEMENTS = {
    "five-method": "five method",
    "TF-IDF-weighted": "TF-IDF weighted",
    "logistic-regression": "logistic regression",
    "zero-shot": "zero shot",
    "parameter-efficient": "parameter efficient",
    "fine-tuning": "fine tuning",
    "retrieval-augmented": "retrieval augmented",
    "PEFT-RAG": "PEFT RAG",
    "post-hoc": "post hoc",
    "validation-selected": "validation selected",
    "class-level": "class level",
    "row-level": "row level",
    "document-title-group": "document title group",
    "document-title": "document title",
    "exact-text": "exact text",
    "answer-only": "answer only",
    "full-sequence": "full sequence",
    "label-list": "label list",
    "training-matched": "training matched",
    "start-token": "start token",
    "beginning-of-sequence": "beginning of sequence",
    "closed-label": "closed label",
    "task-specific": "task specific",
    "retrieval-style": "retrieval style",
    "retrieval-aware": "retrieval aware",
    "inference-time": "inference time",
    "legal-domain": "legal domain",
    "English-language": "English language",
    "model-family": "model family",
    "instruction-tuned": "instruction tuned",
    "output-format": "output format",
    "special-token": "special token",
    "loss-masking": "loss masking",
    "non-conflicting": "nonconflicting",
    "non-conflict": "nonconflict",
    "near-duplicate": "near duplicate",
    "train-test": "training and test",
    "top-ranked": "highest ranked",
    "top-one": "top one",
    "three-neighbour": "three neighbour",
    "twenty-first": "twenty first",
    "data-processing": "data processing",
    "cost-efficiency": "cost efficiency",
    "highest-scoring": "highest scoring",
    "A-to-B": "A to B",
    "Original-to-Corrected": "Original to Corrected",
    "4,096-token": "4,096 token",
    "2,048-token": "2,048 token",
    "12-token": "12 token",
    "20-class": "20 class",
    "A-C": "A to C",
    "Cluster-bootstrap": "Cluster bootstrap",
    "cluster-bootstrap": "cluster bootstrap",
    "Data-focused": "Data focused",
    "data-focused": "data focused",
    "list-valued": "list valued",
    "List-valued": "List valued",
    "Low-Rank": "Low Rank",
    "low-rank": "low rank",
    "Macro-F1": "Macro F1",
    "Mistral-based": "Mistral based",
    "PEFT-adapted": "PEFT adapted",
    "Retrieval-based": "Retrieval based",
    "Sentence-transformer": "Sentence transformer",
    "sentence-transformers": "sentence transformers",
    "Six-method": "Six method",
    "six-method": "six method",
    "Validation-only": "Validation only",
    "validation-only": "validation only",
    "answer-focused": "answer focused",
    "bit-for-bit": "bit for bit",
    "clause-level": "clause level",
    "clause-only": "clause only",
    "clause-type": "clause type",
    "correct-label": "correct label",
    "end-of-sequence": "end of sequence",
    "end-to-end": "end to end",
    "expert-annotated": "expert annotated",
    "fixed-size": "fixed size",
    "four-bit": "four bit",
    "grouped-query": "grouped query",
    "held-out": "held out",
    "in-context": "in context",
    "inference-format": "inference format",
    "inference-prompt": "inference prompt",
    "instruction-capable": "instruction capable",
    "isolated-text": "isolated text",
    "label-to-keyword": "label to keyword",
    "language-model": "language model",
    "leading-label": "leading label",
    "leading-valid-label": "leading valid label",
    "memory-saving": "memory saving",
    "method-specific": "method specific",
    "natural-language": "natural language",
    "near-zero": "near zero",
    "non-canonical": "noncanonical",
    "non-retrieval": "nonretrieval",
    "one-epoch": "one epoch",
    "out-of-domain": "out of domain",
    "p-values": "p values",
    "per-class": "per class",
    "prompt-consistency": "prompt consistency",
    "result-preserving": "result preserving",
    "schema-invalid": "schema invalid",
    "seven-billion-parameter": "seven billion parameter",
    "sliding-window": "sliding window",
    "source-contract": "source contract",
    "strict-parser": "strict parser",
    "test-set": "test set",
    "trade-off": "tradeoff",
    "training-objective": "training objective",
    "training-prompt": "training prompt",
    "two-stage": "two stage",
    "unknown-output": "unknown output",
    "wall-clock": "wall clock",
    "within-group": "within group",
    "2,048 token training sequences": "training sequences of up to 2,048 tokens",
    "12 token generation limit": "generation limit of 12 new tokens",
    "4,096 token input limit": "input limit of 4,096 tokens",
}


def paragraph_index(doc: Document, text: str) -> int:
    for i, paragraph in enumerate(doc.paragraphs):
        if paragraph.text.strip() == text:
            return i
    raise ValueError(f"Paragraph not found: {text}")


def replace_run_text(paragraph, replacements: dict[str, str]) -> None:
    for run in paragraph.runs:
        value = run.text
        for old, new in replacements.items():
            value = value.replace(old, new)
            value = value.replace(old.capitalize(), new.capitalize())
        run.text = value


def set_body_text(paragraph, text: str) -> None:
    clear_paragraph(paragraph)
    paragraph.style = "Thesis Body"
    run = paragraph.add_run(text)
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run.font.size = Pt(12)


def insert_body_before(paragraph, text: str) -> Paragraph:
    element = OxmlElement("w:p")
    paragraph._p.addprevious(element)
    inserted = Paragraph(element, paragraph._parent)
    set_body_text(inserted, text)
    return inserted


def apply_factual_updates(doc: Document) -> None:
    """Bring the uploaded manuscript into line with the frozen audit record."""
    for paragraph in doc.paragraphs:
        if paragraph.text.startswith(
            "The data were derived from version 1 of the Contract Understanding"
        ):
            set_body_text(
                paragraph,
                "The data were derived from version 1 of the Contract Understanding "
                "Atticus Dataset (CUAD), an expert annotated resource for legal "
                "contract review (Hendrycks et al., 2021). CUAD contains 510 "
                "contracts and 41 published label categories. In the project's "
                "master CSV, 40 columns ending in Name-Answer were used to identify "
                "the clause types, after which the list valued spans in the "
                "corresponding clause columns were expanded. This produced 12,984 "
                "clause rows before filtering. The analysis used a fixed set of 20 "
                "clause types, leaving 10,940 observations. Each observation "
                "contained the clause text, its clause type label, and the annotated "
                "Document Name-Answer value, stored in the project as contract_title.",
            )
        elif paragraph.text.startswith(
            "The classical baseline represented each clause with TF-IDF"
        ):
            set_body_text(
                paragraph,
                "The classical baseline represented each clause with TF-IDF weighted "
                "unigrams and bigrams. The vocabulary was limited to 5,000 features. "
                "A logistic regression classifier used balanced class weights, a "
                "maximum of 1,000 iterations, and random seed 42. It was implemented "
                "with scikit-learn (Pedregosa et al., 2011). The method generated one "
                "of the 20 labels directly and could not produce an unknown textual "
                "response.",
            )
        elif paragraph.text.startswith("The adapter used rank 16"):
            set_body_text(
                paragraph,
                "The adapter used rank 16, lora_alpha 32, dropout 0.05, and the "
                "q_proj, k_proj, v_proj, and o_proj target modules. Training used one "
                "epoch, a learning rate of 0.0002, effective batch size 16, seed 42, "
                "and training sequences of up to 2,048 tokens. There were 13,631,488 "
                "trainable parameters, representing approximately 0.1879% of the "
                "7,255,363,584 total parameters.",
            )
        elif paragraph.text.startswith("The RAG method embedded the training clauses"):
            paragraph.text = paragraph.text.replace(
                "(Reimers & Gurevych, 2019)",
                "(Reimers and Gurevych, 2019)",
            )
            paragraph.style = "Thesis Body"
        elif paragraph.text.startswith("Condition C was evaluated once"):
            paragraph.text = paragraph.text.replace(
                "tokenisation, deterministic decoding",
                "tokenisation, input limit of 4,096 tokens, deterministic decoding",
            )
            paragraph.style = "Thesis Body"

    heading = next(
        p for p in doc.paragraphs if p.text.strip() == "3.4.4 Retrieval-augmented generation"
    )
    insert_body_before(
        heading,
        "Original PEFT inference used an input limit of 4,096 tokens and a "
        "generation limit of 12 new tokens. These inference limits differ from "
        "the 2,048 token maximum used to construct training sequences and are "
        "reported separately rather than being treated as one shared truncation setting.",
    )

    heading = next(
        p for p in doc.paragraphs if p.text.strip() == "3.4.5 PEFT-RAG hybrid"
    )
    insert_body_before(
        heading,
        "Within each RAG prompt, every retrieved training clause was capped at "
        "550 tokens. The query clause was capped at 1,500 tokens; for a longer "
        "query, 75% of the retained tokens came from its beginning and 25% from "
        "its end. The assembled prompt was then left truncated to the model's "
        "input limit of 4,096 tokens.",
    )

    heading = next(
        p for p in doc.paragraphs if p.text.strip() == "3.5 Corrected PEFT sensitivity study"
    )
    insert_body_before(
        heading,
        "The hybrid used the same three neighbour retrieval records and the same "
        "550, 1,500 and 4,096 token truncation rules as RAG. Its generation limit "
        "was also 12 new tokens.",
    )


def remove_paragraph(paragraph) -> None:
    element = paragraph._element
    element.getparent().remove(element)


def clear_paragraph(paragraph) -> None:
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def wrap_cached_toc(paragraphs) -> None:
    """Wrap visible TOC entries in an updateable Word TOC field."""
    first = paragraphs[0]
    run = OxmlElement("w:r")
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin.set(qn("w:dirty"), "true")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = ' TOC \\o "1-3" \\h \\z \\u '
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    run.extend([begin, instruction, separate])
    insertion = 1 if first._p.pPr is not None else 0
    first._p.insert(insertion, run)

    final_run = OxmlElement("w:r")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    final_run.append(end)
    paragraphs[-1]._p.append(final_run)


def set_update_fields(doc: Document) -> None:
    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.insert(0, update)
    update.set(qn("w:val"), "true")


def bibliography_entries(path: Path) -> list[str]:
    source = path.read_text(encoding="utf-8")
    source = re.sub(r"^# Bibliography\s*", "", source)
    blocks = [b.strip() for b in re.split(r"\n\s*\n", source) if b.strip()]
    entries = []
    for block in blocks:
        entry = " ".join(line.strip() for line in block.splitlines())
        entries.append(entry)
    return entries


def add_hyperlink(paragraph, text: str, url: str) -> None:
    relation = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relation)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.extend([color, underline])
    run.append(props)
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def set_reference_text(paragraph, text: str) -> None:
    clear_paragraph(paragraph)
    paragraph.style = "Thesis Body"
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.left_indent = Cm(1.25)
    paragraph.paragraph_format.first_line_indent = Cm(-1.25)
    paragraph.paragraph_format.line_spacing = 1.15
    paragraph.paragraph_format.space_after = Pt(6)
    pattern = re.compile(r"\*[^*]+\*|https?://\S+")
    position = 0
    for match in pattern.finditer(text):
        if match.start() > position:
            paragraph.add_run(text[position : match.start()])
        raw = match.group(0)
        if raw.startswith("*"):
            italic = paragraph.add_run(raw[1:-1])
            italic.italic = True
        else:
            trailing = "." if raw.endswith(".") else ""
            url = raw[:-1] if trailing else raw
            add_hyperlink(paragraph, url, url)
            if trailing:
                paragraph.add_run(trailing)
        position = match.end()
    if position < len(text):
        paragraph.add_run(text[position:])
    for run in paragraph.runs:
        run.font.name = "Times New Roman"
        run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Times New Roman")
        run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Times New Roman")
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(0, 0, 0)


def revise(input_path: Path, output_path: Path, bibliography_path: Path) -> None:
    doc = Document(input_path)
    apply_factual_updates(doc)

    # Preserve the user's edited Chapters 1 and 2. Only Chapters 3--7 receive
    # the less hyphenated style treatment.
    chapter_three = paragraph_index(doc, "Chapter 3: Methodology")
    bibliography = paragraph_index(doc, "Bibliography")
    for paragraph in doc.paragraphs[chapter_three:bibliography]:
        replace_run_text(paragraph, STYLE_REPLACEMENTS)
    # The document's tables occur in the methodology, results, title page and
    # AI declaration. Applying the same terminology cleanup to their cell text
    # keeps table captions and entries consistent without touching Chapter 1 or 2.
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    replace_run_text(paragraph, STYLE_REPLACEMENTS)

    # Replace the manually typed contents list with a genuine Word TOC field.
    toc_heading = paragraph_index(doc, "Table of Contents")
    chapter_one = paragraph_index(doc, "Chapter 1: Introduction")
    toc_paragraphs = list(doc.paragraphs[toc_heading + 1 : chapter_one])
    for paragraph in toc_paragraphs:
        replace_run_text(paragraph, STYLE_REPLACEMENTS)
    wrap_cached_toc(toc_paragraphs)
    toc_title = doc.paragraphs[toc_heading]
    toc_title.style = "Normal"  # Exclude the TOC title itself from the TOC.
    toc_title.paragraph_format.space_after = Pt(18)
    toc_title.paragraph_format.keep_with_next = True
    for run in toc_title.runs:
        run.font.name = "Times New Roman"
        run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Times New Roman")
        run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Times New Roman")
        run.font.size = Pt(16)
        run.bold = True
        run.font.color.rgb = RGBColor(0, 0, 0)

    # Update the verified bibliography and apply a standard hanging indent.
    bibliography = paragraph_index(doc, "Bibliography")
    declaration = paragraph_index(doc, "Declaration on the Use of Artificial Intelligence Tools")
    reference_paragraphs = list(doc.paragraphs[bibliography + 1 : declaration])
    entries = bibliography_entries(bibliography_path)
    if len(reference_paragraphs) != len(entries):
        raise ValueError(
            f"Expected {len(reference_paragraphs)} reference paragraphs, found {len(entries)} entries"
        )
    for paragraph, entry in zip(reference_paragraphs, entries, strict=True):
        set_reference_text(paragraph, entry)

    set_update_fields(doc)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("input", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument(
        "--bibliography",
        type=Path,
        default=Path(__file__).resolve().parents[1] / "docs" / "bibliography.md",
    )
    args = parser.parse_args()
    revise(args.input, args.output, args.bibliography)
    print(args.output)


if __name__ == "__main__":
    main()
